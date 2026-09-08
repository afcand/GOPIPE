from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import TakeoffItem

HEADER = ["No", "カテゴリ", "名称", "仕様", "場所", "数量", "単位", "取付高さ FL+", "図面の色(参考)", "ページ", "備考"]


# 数量をどう出したか → 現場が読む日本語。「AIが出した数字」を一列で信用させないため、
# 確かな数（表を読んだ）と当てずっぽう（推定）を必ず別の言葉で書く。
_BASIS_NOTE = {
    "table": "図面の表から（員数欄）",
    # 実測(2026-08-20 ハルキ実図面): 同じ図面を2回かけると計数は 6/11 しか一致しない
    # （排煙口 1→2、VD 2→1、吹出口 1→2）。表由来と同じ顔で並べてはいけない。
    "count": "図面上で計数（要数え直し）",
    "measure": "図面の寸法から計算",
    "estimate": "🔴AI推定（要検算）",
    "none": "数量は未取得（人が記入）",
}


def _note(it: TakeoffItem) -> str:
    """備考列の文言。数量の出所を最優先で示し、低信頼は要確認を添える。"""
    parts: list[str] = []
    # 印字を機械で数えた行に「要数え直し」は付けない。あれは画像認識の計数が
    # 2回かけると動くことへの注意書きで、決定的に数えたものには当てはまらない。
    # 両方並べると、どちらを信じればよいのか分からなくなる。
    if it.source not in ("vector_text", "glyph_count"):
        basis = _BASIS_NOTE.get(it.qty_basis or "")
        if basis:
            parts.append(basis)
    if it.qty_cv is not None:
        if it.source == "cv_count":
            parts.append(f"機械計数{it.qty_cv:g}個を採用(AI読み{(it.qty_vision or 0):g}個)")
        else:
            # qty_cv は「同種記号の図面全体の機械計数」。行（場所ごと）と直接は比べない。
            # 一致/不一致の判定は confidence に反映済み（一致=0.85へ・不一致=0.6以下）。
            parts.append(f"機械計数(図面全体)={it.qty_cv:g}個")
    if it.source == "suppressed_hit":
        parts.append("🔴以前この品目は削除されています（要判断）")
    if it.source == "reconciled":
        parts.append("機器表で数量確定")
    elif it.source == "text_table":
        parts.append("機器表から抽出")
    elif it.source == "legend_count":
        parts.append("凡例から記号カウント")
    elif it.source == "glyph_count":
        # 図形を数えるのも決定的だが、同じ形を使う別記号を巻き込みうる。
        # 名前が「AまたはB」になっている行は、そこを人が決める。
        parts.append("凡例の図形と一致する記号を機械で数えた")
    elif it.source == "vector_text":
        # 印字を機械で数えたので数は動かない。ただし数えたのは「ラベル」であって
        # 部材ではない（1本のダクトに2箇所ラベルがあれば2になる）。ここを
        # 言い換えると嘘になるので、備考でそのまま伝える。
        parts.append("図面の印字を機械で数えた（ラベル箇所数・部材数ではない）")
    if it.confidence < 0.7:
        parts.append(f"要確認(信頼度{it.confidence:.2f})")
    return " / ".join(parts)


def write_excel(
    items: list[TakeoffItem],
    out_path: str | Path,
    *,
    gaps: list | None = None,
    unread: list[tuple[int, str]] | None = None,
    failures: list[str] | None = None,
    refrigerant: list | None = None,
) -> Path:
    """拾い出し表を書く。

    gaps / unread / failures は「この表に出てこないもの」。**必ず別シートに出す。**
    出てこない部材は現場から見れば 0 個に見える。無いのではなく数えられなかった、
    を紙の上で言葉にしないと、そのまま見積になる。
    """
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "拾い出し"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    center = Alignment(horizontal="center", vertical="center")

    ws.append(HEADER)
    for col_idx, _ in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    # カテゴリ順で並べてから No を振る
    sorted_items = sorted(items, key=lambda x: (x.category or "zzz", x.page, x.name))
    for i, it in enumerate(sorted_items, start=1):
        ws.append(
            [
                i,
                it.category or "",
                it.name,
                it.spec or "",
                it.location or "",
                it.quantity,
                it.unit,
                # 図面上でその部材が描かれていた色（機械で実測）。設備図は色で
                # 既存再利用/移設/新設や系統を分ける。色が付いていない＝黒だけの
                # 部材は空欄にする（「その他」に混ぜない）。
                # 取付高さ（Z軸）。高さの違う区間をつなぐ立下りは平面図では点に見え、
                # 延長が丸ごと落ちる。人が気づけるよう列に出す。
                (int(it.level_mm) if it.level_mm is not None else ""),
                (f"{it.color}（{it.color_meaning}）" if it.color and it.color_meaning
                 else (it.color or "")),
                it.page,
                _note(it),
            ]
        )

    # サマリーシート
    ws2 = wb.create_sheet("カテゴリ別集計")
    ws2.append(["カテゴリ", "件数", "合計数量(同一単位のみ)", "代表単位"])
    by_cat: dict[str, list[TakeoffItem]] = defaultdict(list)
    for it in sorted_items:
        by_cat[it.category or "その他"].append(it)
    for cat, rows in by_cat.items():
        units = {r.unit for r in rows}
        rep_unit = next(iter(units)) if len(units) == 1 else "混在"
        total = sum(r.quantity for r in rows) if rep_unit != "混在" else ""
        ws2.append([cat, len(rows), total, rep_unit])
    for col_idx in range(1, 5):
        c = ws2.cell(row=1, column=col_idx)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center

    # 「拾えていないもの」シート。空でも作る＝人が探しに行かなくて済むように。
    ws3 = wb.create_sheet("拾えていないもの")
    ws3.append(["区分", "何が", "なぜ拾えていないか", "どうすればよいか", "ページ"])
    for col_idx in range(1, 6):
        c = ws3.cell(row=1, column=col_idx)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center
    for g in gaps or []:
        ws3.append([
            "この図面にあるが数えられない", g.item, g.reason, g.action,
            "・".join(str(x) for x in g.pages[:20]),
        ])
    for pg, text in unread or []:
        ws3.append([
            "型に当てはめられなかった行", text,
            "ラベルらしい書き方ですが、拾い出しの型に載せられませんでした",
            "人の目で確認してください", str(pg),
        ])
    for f in failures or []:
        ws3.append(["読み取れなかった", "", f, "", ""])
    if not (gaps or unread or failures):
        ws3.append(["", "拾えていないものは検出されませんでした", "", "", ""])
    for col, w in zip("ABCDE", [26, 34, 54, 44, 18]):
        ws3.column_dimensions[col].width = w
    for row in ws3.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    # 図面に刷られた冷媒配管サイズ表。数量ではなく「記号→口径」の読み替え表。
    # 平面図の丸記号は図形で描かれていて機械では拾えないので、人がこれを見て拾う。
    if refrigerant:
        ws4 = wb.create_sheet("冷媒配管サイズ表")
        ws4.append(["記号", "液管", "ガス管", "高低圧ガス管"])
        for col_idx in range(1, 5):
            c = ws4.cell(row=1, column=col_idx)
            c.font = header_font
            c.fill = header_fill
            c.alignment = center
        for r in refrigerant:
            ws4.append([r.symbol, r.liquid, r.gas, r.high_low_gas or "—"])
        ws4.append([])
        ws4.append(["※ 図面に印刷された表です。平面図の冷媒ルートに振られた丸記号を、"
                    "この表で口径に読み替えます。数量ではありません。"])
        for col, w in zip("ABCD", [10, 12, 12, 14]):
            ws4.column_dimensions[col].width = w

    widths = [5, 14, 22, 24, 14, 9, 7, 11, 9, 7, 24]  # 取付高さ・図面の色 を追加
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for i, w in enumerate([16, 8, 18, 12], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    wb.save(out_path)
    return out_path
